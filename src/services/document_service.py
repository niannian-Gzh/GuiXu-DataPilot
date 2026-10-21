"""生成文档产物：PDF / Word / md / txt。

和 `export_service` / `report_service` 同属「产出」类：读数据、落文件，
不改库。区别只在形状——导出是表格，报告是固定版式，这里是**一份文档**。

模型给的是 Markdown（标题、段落、列表、表格），这里负责把它落成四种
格式之一：

- md   原样写
- txt  去掉标记，留文字
- docx 标题 / 段落 / 列表；样式交给 Word 默认模板
- pdf  Markdown → HTML → Playwright，和报告同一条渲染链

**落盘位置统一是 `data/exports/`**，产物栏读的就是这里。
"""

from __future__ import annotations

import html
import re
from pathlib import Path

from config import PROJECT_ROOT
from services import stamped_filename

EXPORT_DIR = PROJECT_ROOT / "data" / "exports"

FORMATS = ("pdf", "docx", "md", "txt")

_PDF_TEMPLATE = """<!doctype html><html><head><meta charset="utf-8"><style>
body { font-family: "Microsoft YaHei", "PingFang SC", sans-serif;
       font-size: 14px; line-height: 1.85; color: #222;
       max-width: 760px; margin: 0 auto; padding: 6mm; }
h1 { font-size: 24px; border-bottom: 2px solid #333; padding-bottom: 8px; }
h2 { font-size: 19px; margin-top: 24px; }
h3 { font-size: 16px; margin-top: 18px; }
table { border-collapse: collapse; width: 100%; margin: 12px 0; }
th, td { border: 1px solid #bbb; padding: 6px 9px; text-align: left; }
th { background: #f3f3f3; }
code { background: #f2f2f2; padding: 1px 4px; border-radius: 3px; }
pre { background: #f6f6f6; padding: 10px; overflow: auto; }
blockquote { border-left: 3px solid #ccc; margin-left: 0;
             padding-left: 12px; color: #555; }
</style></head><body><h1>__TITLE__</h1>__BODY__</body></html>"""


def write(title: str, content: str, fmt: str, to_library: bool = False) -> dict:
    fmt = (fmt or "").lower().strip()
    if fmt not in FORMATS:
        raise ValueError(f"不支持的格式：{fmt or '（空）'}。可用：{' / '.join(FORMATS)}")
    if not content or not content.strip():
        raise ValueError("内容不能为空")

    title = (title or "").strip() or "未命名文档"
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    # unique=True：同一秒连写两份同名文档不能互相覆盖
    path = EXPORT_DIR / stamped_filename(
        title, fmt, max_len=40, fallback="文档", unique=True,
    )

    if fmt == "md":
        path.write_text(content, encoding="utf-8")
    elif fmt == "txt":
        path.write_text(_to_text(content), encoding="utf-8")
    elif fmt == "docx":
        _write_docx(path, content)
    else:
        _write_pdf(path, title, content)

    rel = path.resolve().relative_to(PROJECT_ROOT).as_posix()
    result = {
        "ok": True,
        "file_path": rel,
        "filename": path.name,
        "format": fmt,
        "size": path.stat().st_size,
    }

    # 产物默认不进文档库：它是交付物，不是语料。用户明确说「收进文档库」
    # 才做这一下。入库失败不能连产物一起丢——文件已经写好了，如实报告
    if to_library:
        try:
            from rag import service as rag_service
            result["library"] = rag_service.ingest(path, path.name, keep_source=True)
        except Exception as e:
            result["library_error"] = str(e)

    return result


def _to_text(markdown: str) -> str:
    """Markdown → 纯文本。只去标记，不做排版。"""
    out = []
    for line in markdown.splitlines():
        line = re.sub(r"^#{1,6}\s*", "", line)             # 标题
        line = re.sub(r"^\s*[-*+]\s+", "", line)            # 无序列表
        line = re.sub(r"^\s*\d+[.)]\s+", "", line)          # 有序列表
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)        # 粗体
        line = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", line)
        line = re.sub(r"`(.+?)`", r"\1", line)              # 行内代码
        out.append(line)
    return "\n".join(out).strip() + "\n"


def _table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _table_sep(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def _write_docx(path: Path, markdown: str) -> None:
    from docx import Document

    doc = Document()
    lines = markdown.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        # Markdown 表格：连续的 |...| 行，分隔行跳过。工具描述里写了
        # 「表格都可以」，这里不能把它写成一行行竖线
        if _table_row(line):
            rows = []
            while i < len(lines) and _table_row(lines[i]):
                if not _table_sep(lines[i]):
                    rows.append([
                        c.strip() for c in lines[i].strip().strip("|").split("|")
                    ])
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = table.cell(ri, ci)
                        cell.text = _strip_inline(
                            row[ci] if ci < len(row) else ""
                        )
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            doc.add_heading(_strip_inline(heading.group(2)),
                            level=len(heading.group(1)))
            i += 1
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            doc.add_paragraph(_strip_inline(bullet.group(1)), style="List Bullet")
            i += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            doc.add_paragraph(_strip_inline(numbered.group(1)), style="List Number")
            i += 1
            continue

        doc.add_paragraph(_strip_inline(line))
        i += 1

    doc.save(str(path))


def _strip_inline(text: str) -> str:
    """Word 段落里不认 Markdown 行内标记，去掉即可。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1（\2）", text)
    return text.strip()


def _write_pdf(path: Path, title: str, markdown: str) -> None:
    import markdown as md_lib

    from services.report.render import html_to_pdf

    # 先把 `<` 转义：模型给的正文里若混进 <script> 之类，不能让它
    # 在 Playwright 里执行。只转 `<` 不转 `>`——HTML 标签靠 `<` 成立，
    # 而 `>` 是 Markdown 引用块的行首标记，一起转会把引用吃掉。
    safe = markdown.replace("<", "&lt;")
    body = md_lib.markdown(safe, extensions=["tables", "fenced_code"])
    # 用 replace 不用 format：模板里的 CSS 花括号会被 format 当成占位符
    page = (_PDF_TEMPLATE
            .replace("__TITLE__", html.escape(title))
            .replace("__BODY__", body))
    html_to_pdf(page, path)
