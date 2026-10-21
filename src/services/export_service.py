import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from nl2sql import generate_filter_sql
from execute import execute_sql
from audit import log
from config import PROJECT_ROOT
from services import card_columns, stamped_filename

EXPORT_DIR = PROJECT_ROOT / "data" / "exports"


def run(filter_question: str, ctx) -> str:
    try:
        sql = generate_filter_sql(filter_question, ctx.llm, ctx.schema, ctx.table_name)
        df = execute_sql(sql)
    except Exception as e:
        return json.dumps(
            {"ok": False, "error": f"导出失败：{e}", "error_type": "unknown"},
            ensure_ascii=False,
        )

    if len(df) == 0:
        return json.dumps({
            "ok": True,
            "status": "empty",
            "message": "没有找到符合条件的记录，未生成文件。",
        }, ensure_ascii=False)

    # 列名从**表定义**取（没标中文的回落成字段名）。
    # 旧版这里有一张写死的 EXPORT_COLUMNS——跟确认卡片用的是同一件事，
    # 两处各写一份，加一列就要改两个地方。
    keys, labels = card_columns(ctx)
    rename = dict(zip(keys, labels))
    available = [c for c in df.columns if c in rename]
    df = df[available].rename(columns=rename)

    # 日期字段转字符串
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.strftime("%Y-%m-%d").fillna("")

    # 写文件
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    # 文件名用筛选条件，清洗规则和报告、文档产物共用一份
    filename = stamped_filename(
        filter_question, "xlsx", max_len=30, fallback="导出数据", unique=True,
    )
    file_path = EXPORT_DIR / filename

    df.to_excel(file_path, index=False)

    # 设置列宽，避免日期显示 ##### 
    wb = load_workbook(file_path)
    ws = wb.active
    for i, col_name in enumerate(df.columns, 1):
        values = [str(col_name)] + [str(v) for v in df[col_name]]
        max_len = max(len(v) for v in values) if values else 10
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)
    wb.save(file_path)

    log("export", {
        "trace_id": ctx.trace.id,
        "user_input": ctx.trace.data["question"],
        "file": str(file_path),
        "row_count": len(df),
    })

    return json.dumps({
        "ok": True,
        "status": "ok",
        "file_path": str(file_path.relative_to(PROJECT_ROOT)),
        "row_count": len(df),
        "answer": f"已导出 {len(df)} 条记录到 {file_path.relative_to(PROJECT_ROOT)}",
    }, ensure_ascii=False)